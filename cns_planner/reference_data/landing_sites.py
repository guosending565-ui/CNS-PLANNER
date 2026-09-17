"""Read-only landing-site reference import.

The returned collection is evidence for a human selection workflow.  It is not
the route planner's ``flow.nodes`` collection and does not declare a CRS that
the source workbook itself does not declare.
"""

from __future__ import annotations

import csv
from copy import deepcopy
from hashlib import sha256
from math import asin, cos, radians, sin, sqrt
from pathlib import Path
import re

from ..domain.reference_crs import (
    CRS84, empty_crs_record, is_resolved, normalize_crs_record, unresolved_reason,
)


COLLECTION_ID = "reference-landing-sites"
CRS_STATUS = "pending_confirmation"
LENGTH_SEMANTICS = "geodesic_length_requires_confirmed_source_crs"
_COORDINATE_HINT = re.compile(r"(?:东经|北纬|[EWNS]|[°度])", re.IGNORECASE)
_DMS = re.compile(
    r"(?P<hem>东经|北纬|[EWNS])?\s*[gG]?\s*"
    r"(?P<deg>\d{1,3}(?:\.\d+)?)\s*[°度]\s*"
    r"(?P<minute>\d{1,2}(?:\.\d+)?)\s*['′’分]\s*"
    r"(?P<second>\d{1,2}(?:\.\d+)?)\s*(?:[\"″”秒])?",
    re.IGNORECASE,
)
_DECIMAL_WITH_LABEL = re.compile(
    r"(?P<hem>东经|北纬|[EWNS])\s*(?P<value>-?\d{1,3}(?:\.\d+)?)",
    re.IGNORECASE,
)
_DECIMAL_PAIR = re.compile(
    r"^\s*(?P<lon>-?\d{2,3}(?:\.\d+)?)\s*[,，/\s]+"
    r"(?P<lat>-?\d{1,2}(?:\.\d+)?)\s*$"
)
_HEADER_ALIASES = {
    "serial": ("序号", "编号"),
    "site_type": ("起降设施分类", "起降设施类型", "设施分类", "设施类型", "类型"),
    "region": ("所属县区", "行政区域", "所属区域", "县区", "区域"),
    "location": ("具体位置", "位置", "点位名称", "起降点名称", "名称"),
    "coordinate": ("经纬度信息", "经纬度", "坐标信息", "坐标"),
    "area_m2": ("占地面积", "面积"),
    "completed_time": ("建成时间", "建设时间"),
    "remarks": ("备注",),
}


def empty_reference_landing_sites():
    return {
        "status": "not_calculated",
        "collection_id": COLLECTION_ID,
        "source": None,
        "metadata": {
            "source_type": "real",
            "source_mode": "real",
            "crs_status": CRS_STATUS,
            "crs_status_legacy": CRS_STATUS,
            "planning_integration": "selection_required",
        },
        "crs": empty_crs_record(note=(
            "源表未声明 CRS；坐标仅按源数值临时展示，不得自动假定 WGS84/CGCS2000。"
        )),
        "count": 0,
        "items": [],
        "warnings": [],
    }


def _representation_crs(suffix):
    if suffix in (".geojson", ".json"):
        return {
            "value": CRS84, "status": "declared", "axis_order": "lon_lat",
            "declared_by_format": True,
            "source": {"type": "format_default", "format": "geojson_rfc7946"},
            "evidence": [{
                "type": "rfc7946_format_default", "value": CRS84,
                "note": "RFC 7946 默认只描述 representation；不证明原始 source CRS。",
            }],
        }
    return {
        "value": None, "status": "pending_confirmation", "axis_order": None,
        "declared_by_format": False, "source": None, "evidence": [],
    }


def _axis(hemisphere):
    value = str(hemisphere or "").upper()
    if value in ("E", "W", "东经"):
        return "lon"
    if value in ("N", "S", "北纬"):
        return "lat"
    return None


def _signed(value, hemisphere):
    return -value if str(hemisphere or "").upper() in ("W", "S") else value


def parse_coordinate(raw):
    """Parse common Chinese/E/N decimal and DMS pairs without asserting a CRS."""
    original = "" if raw is None else str(raw).strip()
    warnings = []
    if not original:
        return {"coordinate": None, "quality": "invalid", "warnings": ["missing_coordinate"]}
    estimated = "估" in original
    text = (
        original.replace("\xa0", " ").replace("，", ",").replace("：", ":")
        .replace("‘", "'").replace("’", "'").replace("＇", "'")
        .replace("“", '"').replace("”", '"').replace("＂", '"')
    )
    if re.search(r"(?:^|[^A-Za-z])[gG]\s*\d", text):
        warnings.append("unrecognized_g_prefix")

    values = {}
    unlabeled = []
    matches = list(_DMS.finditer(text))
    for match in matches:
        hemisphere = match.group("hem")
        value = (
            float(match.group("deg"))
            + float(match.group("minute")) / 60.0
            + float(match.group("second")) / 3600.0
        )
        axis = _axis(hemisphere)
        if axis:
            values[axis] = _signed(value, hemisphere)
        else:
            unlabeled.append(value)

    if len(values) + len(unlabeled) < 2:
        labeled = list(_DECIMAL_WITH_LABEL.finditer(text))
        if len(labeled) >= 2:
            values = {
                _axis(item.group("hem")): _signed(float(item.group("value")), item.group("hem"))
                for item in labeled
            }
            unlabeled = []
        else:
            pair = _DECIMAL_PAIR.match(re.sub(r"(?:估|约|东经|北纬|[EN])\s*:?　?", "", text, flags=re.IGNORECASE))
            if pair:
                values = {"lon": float(pair.group("lon")), "lat": float(pair.group("lat"))}
                unlabeled = []

    if unlabeled:
        if "lon" not in values and unlabeled:
            values["lon"] = unlabeled.pop(0)
        if "lat" not in values and unlabeled:
            values["lat"] = unlabeled.pop(0)
        warnings.append("hemisphere_inferred_from_coordinate_order")

    lon, lat = values.get("lon"), values.get("lat")
    if lon is None or lat is None:
        return {"coordinate": None, "quality": "invalid", "warnings": warnings + ["coordinate_parse_failed"]}
    if not (-180 <= lon <= 180 and -90 <= lat <= 90):
        return {"coordinate": None, "quality": "invalid", "warnings": warnings + ["coordinate_out_of_range"]}
    if estimated:
        warnings.append("source_marks_coordinate_estimated")
        quality = "estimated"
    elif "hemisphere_inferred_from_coordinate_order" in warnings:
        quality = "uncertain"
    else:
        quality = "parsed"
    return {"coordinate": [round(lon, 10), round(lat, 10)], "quality": quality, "warnings": warnings}


def _site_type(value, name):
    text = " ".join(str(item or "") for item in (value, name))
    for label in ("公共无人机起降场", "公共起降场", "A1直升机起降场", "直升机场起降点", "直升机起降点", "无人机起降点", "临时起降点", "起降场", "起降点"):
        if label in text:
            return label
    return None


def _normalized_text(value):
    return re.sub(r"[\s\n\r\t（）()：:，,/_-]+", "", str(value or "")).casefold()


def _header_map(values):
    result = {}
    normalized = [_normalized_text(value) for value in values]
    for key, aliases in _HEADER_ALIASES.items():
        for index, value in enumerate(normalized):
            if any(_normalized_text(alias) in value for alias in aliases):
                result[key] = index
                break
    return result if "serial" in result and "coordinate" in result else None


def _looks_like_region(value):
    text = str(value or "").strip()
    return bool(text) and any(text.endswith(suffix) for suffix in ("区", "县", "市", "新区", "管委会", "功能区"))


def _record_from_row(values, source_file, sheet, row_number, headers=None):
    cells = list(values[:8]) + [None] * max(0, 8 - len(values))
    headers = headers or {}
    serial_index = headers.get("serial", 0)
    serial = cells[serial_index] if serial_index < len(cells) else None
    if not isinstance(serial, (int, float)) and not str(serial or "").strip().isdigit():
        return None
    coordinate_index = headers.get("coordinate")
    if coordinate_index is None:
        coordinate_index = next(
        (index for index, value in enumerate(cells) if isinstance(value, str) and _COORDINATE_HINT.search(value)),
        None,
    )
    if coordinate_index is None or coordinate_index >= len(cells):
        return None
    coordinate_raw = cells[coordinate_index]
    parsed = parse_coordinate(coordinate_raw)
    warnings = list(parsed["warnings"])
    location_index = headers.get("location", coordinate_index - 1)
    location = cells[location_index] if 0 <= location_index < len(cells) else None
    region_index, type_index = headers.get("region"), headers.get("site_type")
    region = cells[region_index] if region_index is not None and region_index < len(cells) else None
    explicit_type = cells[type_index] if type_index is not None and type_index < len(cells) else None
    candidates = [
        value for index, value in enumerate(cells[:coordinate_index])
        if index not in (serial_index, location_index) and value not in (None, "")
    ]
    if region in (None, ""):
        region = next((value for value in candidates if _looks_like_region(value)), None)
    if explicit_type in (None, ""):
        explicit_type = next((value for value in candidates if _site_type(value, value)), None)
    name = location or explicit_type
    if explicit_type and "起降" in str(explicit_type) and not location:
        warnings.append("site_name_derived_from_type")
    if not explicit_type:
        warnings.append("site_type_derived_from_name")
    site_type = _site_type(explicit_type, name)
    raw = {f"column_{index + 1}": value for index, value in enumerate(cells) if value is not None}
    attributes = {}
    for key in ("area_m2", "completed_time", "remarks"):
        index = headers.get(key)
        if index is not None and index < len(cells) and cells[index] is not None:
            attributes[key] = cells[index]
    identity = "|".join((
        _normalized_text(region), _normalized_text(site_type), _normalized_text(name),
        _normalized_text(location),
        ",".join(str(value) for value in (parsed["coordinate"] or [])),
        _normalized_text(coordinate_raw) if parsed["coordinate"] is None else "",
        "|".join(f"{key}={_normalized_text(value)}" for key, value in sorted(attributes.items())),
    ))
    return {
        "reference_site_id": None,
        "_identity": identity,
        "name": str(name or location or f"来源行 {row_number}"),
        "region": None if region is None else str(region),
        "site_type": site_type,
        "location": None if location is None else str(location),
        "coordinate_raw": str(coordinate_raw),
        "coordinate": parsed["coordinate"],
        "quality": parsed["quality"],
        "crs_status": CRS_STATUS,
        "warnings": warnings,
        "possible_duplicate": False,
        "duplicate_candidates": [],
        "attributes": attributes,
        "raw": raw,
        "source": {"file": source_file, "sheet": sheet, "row": row_number},
    }


def _haversine_m(left, right):
    lon1, lat1, lon2, lat2 = map(radians, (*left, *right))
    dlon, dlat = lon2 - lon1, lat2 - lat1
    value = sin(dlat / 2) ** 2 + cos(lat1) * cos(lat2) * sin(dlon / 2) ** 2
    return 6371008.8 * 2 * asin(sqrt(value))


def _mark_duplicates(items, threshold_m=75.0):
    for left_index, left in enumerate(items):
        for right in items[left_index + 1:]:
            same_name = left["name"].strip().casefold() == right["name"].strip().casefold()
            close = bool(left["coordinate"] and right["coordinate"] and _haversine_m(left["coordinate"], right["coordinate"]) <= threshold_m)
            if not (same_name or close):
                continue
            for item, other in ((left, right), (right, left)):
                item["possible_duplicate"] = True
                item["duplicate_candidates"].append(other["reference_site_id"])
                if "possible_duplicate" not in item["warnings"]:
                    item["warnings"].append("possible_duplicate")


def _assign_stable_ids(items):
    groups = {}
    for item in items:
        groups.setdefault(item.pop("_identity"), []).append(item)
    for identity, matches in groups.items():
        digest = sha256(identity.encode("utf-8")).hexdigest()[:12].upper()
        for index, item in enumerate(sorted(matches, key=lambda value: json_safe_raw(value["raw"])), 1):
            suffix = f"-{index}" if len(matches) > 1 else ""
            item["reference_site_id"] = f"RLS-{digest}{suffix}"


def json_safe_raw(value):
    return "|".join(f"{key}={value[key]}" for key in sorted(value))


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency contract
        raise ValueError("读取 XLSX 需要安装 geo extra 中的 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            headers = None
            for row_number, values in enumerate(sheet.iter_rows(min_col=1, max_col=8, values_only=True), 1):
                candidate = _header_map(values)
                if candidate:
                    headers = candidate
                    continue
                yield sheet.title, row_number, values, headers
    finally:
        workbook.close()


def _read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        headers = None
        for row_number, values in enumerate(csv.reader(handle), 1):
            candidate = _header_map(values)
            if candidate:
                headers = candidate
                continue
            yield path.stem, row_number, values, headers


def load_reference_landing_sites(path, crs=None):
    """Import landing sites.  ``crs`` may carry a human-confirmed *source* CRS."""

    source_path = Path(str(path or "")).expanduser()
    if not source_path.is_file():
        raise ValueError("参考起降点源路径必须是存在的具体文件")
    suffix = source_path.suffix.lower()
    if suffix == ".et":
        result = empty_reference_landing_sites()
        result.update({
            "status": "requires_xlsx_or_csv_conversion",
            "source": {"file": source_path.name, "format": "et"},
            "warnings": ["requires_xlsx_or_csv_conversion"],
        })
        return result
    if suffix not in (".xlsx", ".csv"):
        raise ValueError("参考起降点仅支持 XLSX 或 CSV；ET 需先转换")
    crs_record = normalize_crs_record(crs) if crs is not None else empty_crs_record()
    crs_record["representation_crs"] = _representation_crs(suffix)
    source_resolved = is_resolved(crs_record, role="source_crs")
    reader = _read_xlsx(source_path) if suffix == ".xlsx" else _read_csv(source_path)
    items = []
    for sheet, row_number, values, headers in reader:
        item = _record_from_row(values, source_path.name, sheet, row_number, headers)
        if item:
            item["crs"] = deepcopy(crs_record)
            item["metric_measurement_status"] = "enabled" if source_resolved else "disabled_unresolved_source_crs"
            items.append(item)
    _assign_stable_ids(items)
    _mark_duplicates(items)
    quality_counts = {name: sum(item["quality"] == name for item in items) for name in ("parsed", "estimated", "uncertain", "invalid")}
    warnings = ["source_crs_pending_confirmation"]
    if crs_record["representation_crs"]["declared_by_format"]:
        warnings.append("representation_crs_declared_by_format_not_source_crs")
    if not source_resolved:
        warnings.append("metric_measurement_disabled_unresolved_source_crs")
    return {
        "status": "passed" if items else "missing_data",
        "collection_id": COLLECTION_ID,
        "source": {"file": source_path.name, "format": suffix.lstrip(".")},
        "metadata": {
            "source_type": "real", "source_mode": "real",
            "crs_status": CRS_STATUS,
            "crs_status_legacy": CRS_STATUS,
            "planning_integration": "selection_required",
            "coordinate_order": "lon_lat",
            "quality_counts": quality_counts,
            "possible_duplicate_count": sum(item["possible_duplicate"] for item in items),
            "metric_measurement_status": "enabled" if source_resolved else "disabled_unresolved_source_crs",
            "length_semantics": LENGTH_SEMANTICS,
        },
        "crs": crs_record,
        "count": len(items),
        "items": items,
        "warnings": warnings,
    }


def backfill_reference_landing_sites(collection):
    """Additive CRS backfill for landing-site collections saved before the split.

    Idempotent, and an untouched empty collection is returned exactly as
    :func:`empty_reference_landing_sites` wrote it so project round-trips stay stable.
    """

    if not isinstance(collection, dict) or not collection:
        return empty_reference_landing_sites()
    if not collection.get("items") and collection.get("status") in (None, "not_calculated"):
        return empty_reference_landing_sites()
    result = deepcopy(collection)
    result.setdefault("collection_id", COLLECTION_ID)
    result.setdefault("metadata", {})
    result["metadata"].setdefault("crs_status_legacy", result["metadata"].get("crs_status", CRS_STATUS))
    crs = normalize_crs_record(result.get("crs"))
    if result.get("crs") is None and result["metadata"].get("crs_status"):
        crs = normalize_crs_record({"crs_status": result["metadata"]["crs_status"]})
    result["crs"] = crs
    source_resolved = is_resolved(crs, role="source_crs")
    result["metadata"]["metric_measurement_status"] = (
        "enabled" if source_resolved else "disabled_unresolved_source_crs"
    )
    result["metadata"].setdefault("length_semantics", LENGTH_SEMANTICS)
    warnings = list(result.get("warnings") or [])
    if not source_resolved and "metric_measurement_disabled_unresolved_source_crs" not in warnings:
        warnings.append("metric_measurement_disabled_unresolved_source_crs")
    result["warnings"] = warnings
    for item in result.get("items") or []:
        if not isinstance(item, dict):
            continue
        item.setdefault("crs", deepcopy(crs))
        item["metric_measurement_status"] = (
            "enabled" if source_resolved else "disabled_unresolved_source_crs"
        )
    return result


def reference_crs_status(collection):
    """Compact readiness view used by the data-readiness panel and evidence pack."""

    crs = (collection or {}).get("crs") or empty_crs_record()
    return {
        "source_crs": deepcopy(crs.get("source_crs")),
        "representation_crs": deepcopy(crs.get("representation_crs")),
        "source_resolved": is_resolved(crs, role="source_crs"),
        "representation_resolved": is_resolved(crs, role="representation_crs"),
        "unresolved_reason": unresolved_reason(crs, role="source_crs"),
    }
