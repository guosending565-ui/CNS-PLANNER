"""Reference-only route and waypoint imports from converted source files."""

from __future__ import annotations

import csv
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re

from ..domain.geodesy import distance_m
from ..domain.reference_crs import CRS84, empty_crs_record, is_resolved, normalize_crs_record, unresolved_reason
from .landing_sites import CRS_STATUS, parse_coordinate


COLLECTION_ID = "reference-routes"
#: Length/similarity stay disabled until the *source* CRS is confirmed by a human.
LENGTH_SEMANTICS = "geodesic_length_requires_confirmed_source_crs"
#: 当前的支持列名（``_key`` 会先去掉下划线/空格/全角标点再做匹配，因此
#: ``route_id`` 与 ``routeId`` 命中同一条目）。新增的 ``route_id`` / ``point_order``
#: 等别名只做兼容读取，不改变分组与排序算法。
_ALIASES = {
    "route_number": ("航线编号", "航路编号", "航线号", "route_number", "route_no", "route_id", "route_code"),
    "route_name": ("航线名称", "航路名称", "route_name"),
    "category": ("航线类别", "航线类型", "航路类别", "category"),
    "sequence": (
        "点序号", "点位序号", "航点序号", "航路点序号", "点序", "序号",
        "sequence", "seq", "point_order", "point_seq", "point_index", "order",
    ),
    "point_name": ("点位名称", "航点名称", "航路点名称", "点名称", "名称", "point_name"),
    "point_type": ("点位类型", "航点类型", "航路点类型", "点类型", "point_type"),
    "coordinate": ("经纬度信息", "经纬度", "坐标信息", "坐标", "coordinate"),
    "longitude": ("经度", "经度E", "E经度", "longitude", "lon", "lng", "x"),
    "latitude": ("纬度", "纬度N", "N纬度", "latitude", "lat", "y"),
    # 源文件自带的坐标系声明。这里只负责转述文件写了什么，绝不做任何坐标系推断；
    # 只有 ``crs_confirmed`` 明确为真时，应用层才会用 pyproj 校验并自动确认。
    "source_crs": ("源坐标系", "源CRS", "源坐标系统", "source_crs", "epsg", "srid"),
    "crs_confirmed": ("CRS已确认", "坐标系已确认", "crs_confirmed", "source_crs_confirmed"),
}

_TRUTHY = ("true", "1", "yes", "y", "是", "已确认", "confirmed")
_FALSY = ("false", "0", "no", "n", "否", "未确认", "unconfirmed", "pending")


def empty_reference_routes():
    return {
        "status": "not_calculated",
        "collection_id": COLLECTION_ID,
        "source": None,
        "metadata": {
            "source_type": "real", "source_mode": "real",
            "crs_status": CRS_STATUS, "usage": "reference_only",
            "planning_integration": "never_automatic",
            "crs_status_legacy": CRS_STATUS,
        },
        "crs": empty_crs_record(note=(
            "源文件未声明 CRS；CSV/XLSX 不得被自动假定为 WGS84/CGCS2000。"
            "source_crs 未确认前不输出正式 length_m，也不输出米制几何相似度。"
        )),
        "count": 0,
        "point_count": 0,
        "items": [],
        "points": [],
        "warnings": [],
        # 数据源身份/导入状态摘要与源文件自身的坐标系声明。空集合保持为 ``None``，
        # 使"未配置"与"已导入"在项目状态里结构对称。
        "declared_source_crs": None,
        "data_source": None,
    }


def _format_crs(suffix):
    """Representation CRS implied by the *format* alone (never the source CRS)."""

    if suffix in (".geojson", ".json"):
        return {
            "value": CRS84,
            "status": "declared",
            "axis_order": "lon_lat",
            "declared_by_format": True,
            "source": {"type": "format_default", "format": "geojson_rfc7946"},
            "evidence": [{
                "type": "rfc7946_format_default",
                "value": CRS84,
                "note": (
                    "RFC 7946 GeoJSON 规定坐标使用 OGC:CRS84；这只描述 representation，"
                    "不能证明原始测量/调查 CRS。"
                ),
            }],
        }
    return {
        "value": None, "status": "pending_confirmation", "axis_order": None,
        "declared_by_format": False, "source": None, "evidence": [],
    }


def _key(value):
    return re.sub(r"[\s\n\r\t（）()：:，,/_-]+", "", str(value or "")).casefold()


def _field(row, name):
    indexed = {_key(key): value for key, value in row.items()}
    for alias in _ALIASES[name]:
        value = indexed.get(_key(alias))
        if value not in (None, ""):
            return value
    return None


def _number(value):
    try:
        text = str(value).strip()
        match = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*[°度]?", text)
        number = float(match.group(1) if match else text)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number


def _boolean(value):
    """源文件里的确认标记；无法识别时返回 ``None``（绝不猜测）。"""

    if isinstance(value, bool):
        return value
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().casefold()
    if text in _TRUTHY:
        return True
    if text in _FALSY:
        return False
    return None


def _coordinate(row):
    lon_raw, lat_raw = _field(row, "longitude"), _field(row, "latitude")
    lon, lat = _number(lon_raw), _number(lat_raw)
    if lon is not None and lat is not None:
        if -180 <= lon <= 180 and -90 <= lat <= 90:
            return {"coordinate": [lon, lat], "quality": "parsed", "warnings": []}
        return {"coordinate": None, "quality": "invalid", "warnings": ["coordinate_out_of_range"]}
    combined = _field(row, "coordinate")
    if combined in (None, "") and lon_raw not in (None, "") and lat_raw not in (None, ""):
        combined = f"E{lon_raw} N{lat_raw}"
    return parse_coordinate(combined)


def _read_csv(path):
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise ValueError("CSV 缺少表头")
                return [(row_number, dict(row)) for row_number, row in enumerate(reader, 2)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError("参考航线 CSV 编码无法识别") from last_error


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("读取 XLSX 需要安装 geo extra 中的 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        for sheet in workbook.worksheets:
            headers = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                values = list(values)
                keys = {_key(value) for value in values if value not in (None, "")}
                route_aliases = {_key(value) for value in _ALIASES["route_number"]}
                sequence_aliases = {_key(value) for value in _ALIASES["sequence"]}
                if keys & route_aliases and keys & sequence_aliases:
                    headers = [str(value or "") for value in values]
                    continue
                if headers and any(value not in (None, "") for value in values):
                    rows.append((sheet.title, row_number, dict(zip(headers, values))))
    finally:
        workbook.close()
    return rows


def _geojson_rows(path):
    document = json.loads(path.read_text(encoding="utf-8-sig"))
    features = document.get("features") if isinstance(document, dict) else None
    if not isinstance(features, list):
        raise ValueError("参考航线 GeoJSON 必须是 FeatureCollection")
    rows = []
    for feature_number, feature in enumerate(features, 1):
        properties = dict((feature or {}).get("properties") or {})
        geometry = (feature or {}).get("geometry") or {}
        coordinates = geometry.get("coordinates") or []
        if geometry.get("type") == "Point":
            properties["longitude"], properties["latitude"] = coordinates[:2]
            rows.append((feature_number, properties))
        elif geometry.get("type") == "LineString":
            route_number = _field(properties, "route_number") or feature_number
            for sequence, coordinate in enumerate(coordinates, 1):
                row = dict(properties)
                row.update({"route_number": route_number, "sequence": sequence,
                            "longitude": coordinate[0], "latitude": coordinate[1]})
                rows.append((feature_number, row))
    return rows


def _declaration(rows):
    """转述源文件自身的 ``source_crs`` / ``crs_confirmed`` 声明。

    这里只回答"文件写了什么"：绝不根据坐标数值、文件名或格式推断坐标系，
    也不把声明直接升级为已确认记录——那一步由应用层用 pyproj 校验后完成。
    未声明坐标系时返回 ``None``。
    """

    value, confirmed = None, None
    for _sheet, _row_number, row in rows:
        current = _field(row, "source_crs")
        if value is None and current not in (None, ""):
            value = str(current).strip()
        flag = _boolean(_field(row, "crs_confirmed"))
        if flag is not None and confirmed is None:
            confirmed = flag
    if not value:
        return None
    return {
        "value": value,
        "confirmed": confirmed is True,
        "columns": ["source_crs", "crs_confirmed"],
        "evidence": [{
            "type": "source_file_declaration",
            "value": value,
            "confirmed": confirmed is True,
            "note": (
                "源文件 source_crs 列声明该坐标系；crs_confirmed 列"
                f"{'为真' if confirmed is True else '未为真'}。"
                "该记录只转述文件内容，仍须 pyproj 校验后才作为已确认坐标系。"
            ),
        }],
    }


def declared_source_crs(path):
    """读取任意受支持参考航线文件的坐标系声明；不可读或格式不支持时返回 ``None``。"""

    source_path = Path(str(path or "")).expanduser()
    if not source_path.is_file():
        return None
    suffix = source_path.suffix.lower()
    try:
        if suffix == ".csv":
            rows = [(source_path.stem, row, values) for row, values in _read_csv(source_path)]
        elif suffix == ".xlsx":
            rows = _read_xlsx(source_path)
        elif suffix in (".geojson", ".json"):
            rows = [(source_path.stem, row, values) for row, values in _geojson_rows(source_path)]
        else:
            return None
    except (ValueError, OSError):
        return None
    return _declaration(rows)


def _route_id(route_number):
    return "RLR-" + sha256(_key(route_number).encode("utf-8")).hexdigest()[:12].upper()


def _point_id(route_number, sequence, name, point_type, coordinate):
    identity = "|".join((
        _key(route_number), str(sequence), _key(name), _key(point_type),
        ",".join(str(value) for value in coordinate or []),
    ))
    return "RLP-" + sha256(identity.encode("utf-8")).hexdigest()[:12].upper()


def backfill_reference_routes(collection):
    """Additive backfill for projects saved before the CRS split.

    Legacy records kept only ``crs_status`` and an unconditional ``length_m``.  The
    legacy status becomes visible evidence, the metric length is re-derived from the
    (unresolved) CRS record — which means it becomes ``None`` — and the original value
    is preserved as ``legacy_length_m`` so nothing is silently lost.

    The function is idempotent *and* leaves an untouched empty collection exactly as
    :func:`empty_reference_routes` wrote it, so a project round-trip stays stable.
    """

    if not isinstance(collection, dict):
        return empty_reference_routes()
    if not collection:
        return empty_reference_routes()
    if not collection.get("items") and not collection.get("points") and (
        collection.get("status") in (None, "not_calculated")
    ):
        return empty_reference_routes()
    result = deepcopy(collection)
    result.setdefault("collection_id", COLLECTION_ID)
    result.setdefault("metadata", {})
    result["metadata"].setdefault("crs_status_legacy", result["metadata"].get("crs_status", CRS_STATUS))
    crs = normalize_crs_record(result.get("crs"))
    if result.get("crs") is None and result["metadata"].get("crs_status"):
        crs = normalize_crs_record({"crs_status": result["metadata"]["crs_status"]})
    result["crs"] = crs
    result["metadata"]["length_semantics"] = LENGTH_SEMANTICS
    source_resolved = is_resolved(crs, role="source_crs")
    result["metadata"]["metric_measurement_status"] = (
        "enabled" if source_resolved else "disabled_unresolved_source_crs"
    )
    warnings = list(result.get("warnings") or [])
    if not source_resolved and "length_m_and_metric_similarity_disabled" not in warnings:
        warnings.append("length_m_and_metric_similarity_disabled")
    result["warnings"] = warnings
    for index, route in enumerate(result.get("items") or []):
        if not isinstance(route, dict):
            continue
        route.setdefault("crs", deepcopy(crs))
        route.setdefault("source_numeric_path", deepcopy(route.get("path") or []))
        route.setdefault("length_semantics", LENGTH_SEMANTICS)
        route.setdefault("length_unresolved_reason", unresolved_reason(crs, role="source_crs"))
        route["metric_geometry_similarity_available"] = source_resolved
        if not source_resolved:
            if route.get("length_m") is not None:
                route.setdefault("legacy_length_m", route["length_m"])
            route["length_m"] = None
            route["length_status"] = "blocked_unresolved_source_crs"
        else:
            route["length_status"] = "passed"
        result["items"][index] = route
    for point in result.get("points") or []:
        if isinstance(point, dict):
            point.setdefault("crs", deepcopy(crs))
    # 数据源摘要的加法式回填：旧项目没有该字段时不虚构路径，只保证结构一致；
    # 已有摘要的项目按当前集合内容刷新计数与坐标系状态。
    result.setdefault("declared_source_crs", None)
    source_state = result.get("data_source")
    if isinstance(source_state, dict):
        source_state.setdefault("imported", bool(result.get("items")))
        source_state["route_count"] = len(result.get("items") or [])
        source_state["point_count"] = len(result.get("points") or [])
        source_state["crs"] = crs["source_crs"].get("value") or source_state.get("crs")
        source_state["crs_confirmed"] = source_resolved
    else:
        result["data_source"] = None
    return result


def _build(rows, source_path, crs):
    grouped = {}
    invalid_rows = 0
    last_route_number = None
    last_route_name = None
    last_category = None
    source_resolved = is_resolved(crs, role="source_crs")
    for sheet, row_number, row in rows:
        route_number = _field(row, "route_number")
        sequence = _number(_field(row, "sequence"))
        parsed = _coordinate(row)
        if route_number not in (None, ""):
            last_route_number = route_number
            last_route_name = _field(row, "route_name")
            last_category = _field(row, "category")
        elif sequence is not None:
            route_number = last_route_number
        if route_number in (None, "") or sequence is None:
            invalid_rows += 1
            continue
        route_number = str(route_number).strip()
        point_name = str(_field(row, "point_name") or f"点 {sequence}").strip()
        point_type = _field(row, "point_type")
        route_id = _route_id(route_number)
        point = {
            "reference_route_point_id": _point_id(
                route_number, sequence, point_name, point_type, parsed["coordinate"]
            ),
            "route_id": route_id,
            "route_number": route_number,
            "sequence": sequence,
            "name": point_name,
            "type": None if point_type in (None, "") else str(point_type),
            "coordinate": parsed["coordinate"],
            "quality": parsed["quality"],
            "source": {"file": source_path.name, "sheet": sheet, "row": row_number},
            "crs_status": CRS_STATUS,
            "crs": deepcopy(crs),
            "position": None,
            "warnings": list(parsed["warnings"]),
        }
        grouped.setdefault(route_number, {"rows": [], "points": [], "names": [], "categories": []})
        grouped[route_number]["rows"].append(row_number)
        grouped[route_number]["points"].append(point)
        grouped[route_number]["names"].append(_field(row, "route_name") or last_route_name)
        grouped[route_number]["categories"].append(_field(row, "category") or last_category)

    routes, points = [], []
    for route_number in sorted(grouped, key=lambda value: (_number(value) is None, _number(value) or 0, value)):
        group = grouped[route_number]
        ordered = sorted(group["points"], key=lambda item: (float(item["sequence"]), item["reference_route_point_id"]))
        for index, point in enumerate(ordered):
            point["position"] = "endpoint" if index in (0, len(ordered) - 1) else "intermediate"
        path = [point["coordinate"] for point in ordered if point["coordinate"]]
        warnings = []
        if len(path) != len(ordered):
            warnings.append("invalid_route_point_coordinate")
        if len(path) < 2:
            warnings.append("insufficient_valid_points_for_polyline")
        name = next((str(value).strip() for value in group["names"] if value not in (None, "")), f"航线 {route_number}")
        category = next((str(value).strip() for value in group["categories"] if value not in (None, "")), None)
        route_id = _route_id(route_number)
        for point in ordered:
            point["route_name"] = name
            point["category"] = category
        routes.append({
            "reference_route_id": route_id,
            "route_number": route_number,
            "name": name,
            "category": category,
            "ordered_points": [dict(point) for point in ordered],
            "ordered_point_ids": [point["reference_route_point_id"] for point in ordered],
            "path": path,
            # A metric length is only meaningful against a resolved source CRS.
            "length_m": (
                sum(distance_m(left, right) for left, right in zip(path, path[1:]))
                if source_resolved else None
            ),
            "length_status": "passed" if source_resolved else "blocked_unresolved_source_crs",
            "length_semantics": LENGTH_SEMANTICS,
            "length_unresolved_reason": unresolved_reason(crs, role="source_crs"),
            "metric_geometry_similarity_available": source_resolved,
            "source_numeric_path": deepcopy(path),
            "source": {"file": source_path.name, "rows": sorted(group["rows"])},
            "provenance": {"group_by": "route_number", "order_by": "sequence"},
            "crs_status": CRS_STATUS,
            "crs": deepcopy(crs),
            "warnings": warnings,
            "usage": "reference_only",
        })
        points.extend(ordered)
    return routes, points, invalid_rows


def load_reference_routes(path, crs=None):
    """Import reference routes.  ``crs`` may carry a *human-confirmed* source CRS.

    Absent explicit confirmation the source CRS stays ``pending_confirmation`` and no
    formal ``length_m`` / metric geometry similarity is produced — the coordinates are
    kept as ``source_numeric_path`` for temporary display only.
    """

    source_path = Path(str(path or "")).expanduser()
    if not source_path.is_file():
        raise ValueError("参考航线源路径必须是存在的具体文件")
    suffix = source_path.suffix.lower()
    if suffix == ".et":
        result = empty_reference_routes()
        result.update({
            "status": "requires_xlsx_or_csv_conversion",
            "source": {"file": source_path.name, "format": "et"},
            "warnings": ["requires_xlsx_or_csv_conversion"],
        })
        return result
    representation = _format_crs(suffix)
    crs_record = normalize_crs_record(crs) if crs is not None else empty_crs_record()
    crs_record["representation_crs"] = representation
    crs_record.setdefault("note", (
        "source_crs 未确认前不输出正式 length_m，也不输出米制几何相似度；"
        "representation_crs 仅描述当前内存/渲染解释。"
    ))
    if suffix == ".csv":
        rows = [(source_path.stem, row, values) for row, values in _read_csv(source_path)]
    elif suffix == ".xlsx":
        rows = _read_xlsx(source_path)
    elif suffix in (".geojson", ".json"):
        rows = [(source_path.stem, row, values) for row, values in _geojson_rows(source_path)]
    else:
        raise ValueError("参考航线仅支持 CSV/XLSX/GeoJSON；ET 需先转换")
    routes, points, invalid_rows = _build(rows, source_path, crs_record)
    declaration = _declaration(rows)
    warnings = ["source_crs_pending_confirmation"]
    if representation["declared_by_format"]:
        warnings.append("representation_crs_declared_by_format_not_source_crs")
    if not is_resolved(crs_record, role="source_crs"):
        warnings.append("length_m_and_metric_similarity_disabled")
    return {
        "status": "passed" if routes else "missing_data",
        "collection_id": COLLECTION_ID,
        "source": {"file": source_path.name, "format": suffix.lstrip(".")},
        "metadata": {
            "source_type": "real", "source_mode": "real",
            "crs_status": CRS_STATUS, "usage": "reference_only",
            "planning_integration": "never_automatic",
            "group_by": "route_number", "order_by": "sequence",
            "invalid_row_count": invalid_rows,
            "crs_status_legacy": CRS_STATUS,
            "length_semantics": LENGTH_SEMANTICS,
            "metric_measurement_status": (
                "enabled" if is_resolved(crs_record, role="source_crs") else "disabled_unresolved_source_crs"
            ),
        },
        "crs": crs_record,
        "count": len(routes),
        "point_count": len(points),
        "items": routes,
        "points": points,
        "warnings": warnings,
        # 数据源身份与导入状态的显式摘要：项目保存后据此恢复业务航线，
        # 无需用户重新挑选文件或重新确认坐标系。
        "declared_source_crs": declaration,
        "data_source": {
            "path": str(source_path),
            "file_name": source_path.name,
            "format": suffix.lstrip("."),
            "crs": (declaration or {}).get("value"),
            "crs_confirmed": bool((declaration or {}).get("confirmed")),
            "imported": False,
            "route_count": len(routes),
            "point_count": len(points),
        },
    }
