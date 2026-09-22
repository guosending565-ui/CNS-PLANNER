"""Read-only real communication-tower import.

本模块只做一件事：把外部铁塔站址表读成 :mod:`cns_planner.domain.towers` 契约下的
只读集合，并让**每一条记录都能回到原始文件的具体 sheet 与行号**。

它不声明源文件没有声明的 CRS，不生成覆盖，不把站址转成 CNS 设备。
"""

from __future__ import annotations

import csv
import json
import re
from copy import deepcopy
from pathlib import Path

from ..domain.reference_crs import CRS84, empty_crs_record, normalize_crs_record
from ..domain.towers import (
    COLLECTION_ID, SOURCE_ATTRIBUTE_FIELDS, TOWER_SCHEMA_VERSION, empty_towers,
    normalize_tower, validate_coordinate,
)

#: 单位后缀与括号在匹配表头时被去掉，因此 "海拔高度(m)" 与 "海拔高度" 等价。
_UNIT = re.compile(r"[（(][^）)]*[）)]")
_NON_WORD = re.compile(r"[\s_\-/·、,，:：]+")

_HEADER_ALIASES = {
    "tower_id": (
        "所属站址编码", "站址编码", "铁塔编码", "站点编码", "站址编号", "资源编码",
        "tower_id", "site_id", "site_code", "id",
    ),
    "name": ("站址名称", "站点名称", "铁塔名称", "站名", "名称", "name", "site_name"),
    "longitude": ("经度", "东经", "lon", "lng", "longitude", "x"),
    "latitude": ("纬度", "北纬", "lat", "latitude", "y"),
    "district": ("区域", "所属区域", "行政区", "区县", "district", "region"),
    "site_type": ("铁塔细分类型", "铁塔类型", "塔型", "类型", "site_type", "tower_type"),
    "elevation_m": ("海拔高度", "海拔", "高程", "elevation", "elevation_m", "altitude"),
    "height_m": ("塔身高度", "塔高", "铁塔高度", "挂高", "height", "height_m", "tower_height"),
    "operator": ("运营商", "产权单位", "所属运营商", "operator", "carrier"),
    "address": ("地址", "详细地址", "address"),
    "remarks": ("备注", "说明", "remarks", "note"),
}

#: 认表头时的最小充分条件：至少同时出现 ID 列与经/纬度列。
_IDENTITY_HEADERS = ("tower_id",)
_COORDINATE_HEADERS = ("longitude", "latitude")

#: 近重复判定阈值（度）。仅用于**报告**同一位置存在多条记录，绝不自动合并。
NEAR_DUPLICATE_TOLERANCE_DEG = 1e-5


def _key(value):
    text = _UNIT.sub("", str(value or "")).strip().lower()
    return _NON_WORD.sub("", text)


_KEY_LOOKUP = {}
for _field, _aliases in _HEADER_ALIASES.items():
    for _alias in _aliases:
        _KEY_LOOKUP[_key(_alias)] = _field


def _field_from_key(value):
    return _KEY_LOOKUP.get(_key(value))


def _identifier(value):
    """Keep tower IDs as stable text; never let a float exponent destroy an ID."""

    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    text = str(value).strip()
    return text or None


def _text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _number(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _representation_crs(suffix):
    if suffix in (".geojson", ".json"):
        return {
            "value": CRS84, "status": "declared", "axis_order": "lon_lat",
            "declared_by_format": True,
            "source": {"type": "format_default", "format": "geojson_rfc7946"},
            "evidence": [{
                "type": "rfc7946_format_default", "value": CRS84,
                "note": "RFC 7946 默认只描述 representation；不证明原始 source CRS。",
            }],
        }
    return {
        "value": None, "status": "pending_confirmation", "axis_order": None,
        "declared_by_format": False, "source": None, "evidence": [],
    }


def _mapped_row(raw):
    """Map one raw source row onto contract fields via the alias table."""

    mapped, unknown = {}, {}
    for column, value in (raw or {}).items():
        field = _field_from_key(column)
        if field and mapped.get(field) in (None, "") and value not in (None, ""):
            mapped[field] = value
        elif field is None:
            unknown[str(column)] = _text(value)
    return mapped, unknown


def _rows_to_items(rows, *, source_file):
    items, invalid = [], []
    for sheet, row_number, raw in rows:
        mapped, unknown = _mapped_row(raw)
        reference = {
            "type": "file_row",
            "file": source_file,
            "file_name": Path(source_file).name,
            "sheet": sheet,
            "row": row_number,
        }
        candidate = {
            "tower_id": _identifier(mapped.get("tower_id")),
            "name": _text(mapped.get("name")),
            "longitude": _number(mapped.get("longitude")),
            "latitude": _number(mapped.get("latitude")),
            "district": _text(mapped.get("district")),
            "site_type": _text(mapped.get("site_type")),
            "elevation_m": _number(mapped.get("elevation_m")),
            "height_m": _number(mapped.get("height_m")),
            "operator": _text(mapped.get("operator")),
            "address": _text(mapped.get("address")),
            "remarks": _text(mapped.get("remarks")),
            "source": reference,
            "evidence": [dict(reference)],
            "raw_attributes": unknown,
        }
        if not candidate["tower_id"]:
            invalid.append({"sheet": sheet, "row": row_number, "reason": "missing_tower_id"})
            continue
        if candidate["longitude"] is None or candidate["latitude"] is None:
            invalid.append({"sheet": sheet, "row": row_number, "reason": "missing_coordinate"})
            continue
        check = validate_coordinate([candidate["longitude"], candidate["latitude"]])
        if not check["valid"]:
            invalid.append({"sheet": sheet, "row": row_number, "reason": check["reason"]})
            continue
        items.append(candidate)
    return items, invalid


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("读取 XLSX 需要安装 geo extra 中的 openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        for sheet in workbook.worksheets:
            headers = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                values = list(values)
                fields = {_field_from_key(value) for value in values if value not in (None, "")}
                fields.discard(None)
                if headers is None and (
                    set(_IDENTITY_HEADERS) <= fields and set(_COORDINATE_HEADERS) <= fields
                ):
                    headers = {_field_from_key(value): index for index, value in enumerate(values)}
                    continue
                if headers and any(value not in (None, "") for value in values):
                    rows.append((sheet.title, row_number, {
                        name: (values[index] if index < len(values) else None)
                        for name, index in headers.items()
                    }))
    finally:
        workbook.close()
    return rows


def _read_csv(path):
    last_error = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise ValueError("铁塔 CSV 缺少表头")
                return [(None, number, dict(row)) for number, row in enumerate(reader, 2)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError("铁塔 CSV 编码无法识别") from last_error


def _read_geojson(path):
    document = json.loads(path.read_text(encoding="utf-8-sig"))
    features = document.get("features") if isinstance(document, dict) else None
    if not isinstance(features, list):
        raise ValueError("铁塔 GeoJSON 必须是 FeatureCollection")
    rows = []
    for number, feature in enumerate(features, 1):
        properties = dict((feature or {}).get("properties") or {})
        geometry = (feature or {}).get("geometry") or {}
        if geometry.get("type") != "Point":
            raise ValueError("铁塔 GeoJSON 仅支持 Point 要素")
        coordinates = geometry.get("coordinates") or []
        if len(coordinates) >= 2:
            properties.setdefault("longitude", coordinates[0])
            properties.setdefault("latitude", coordinates[1])
        rows.append((None, number, properties))
    return rows


def _read_rows(path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in (".geojson", ".json"):
        return _read_geojson(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise ValueError(f"铁塔数据暂不支持 {suffix or '无扩展名'}；支持 .xlsx / .csv / .geojson")


def duplicate_report(items):
    """Group records that share a location; never merges them."""

    groups = {}
    for item in items:
        key = (round(float(item["longitude"]), 5), round(float(item["latitude"]), 5))
        groups.setdefault(key, []).append(item["tower_id"])
    duplicates = [
        {"longitude": key[0], "latitude": key[1], "tower_ids": ids, "count": len(ids)}
        for key, ids in sorted(groups.items()) if len(ids) > 1
    ]
    return duplicates


def load_towers(path, crs=None, *, source_crs=None, crs_confirmed=False,
                crs_evidence=None, crs_source=None):
    """Load a tower table read-only.

    ``crs`` may carry an explicit normalized CRS record.  ``source_crs`` /
    ``crs_confirmed`` are the human-confirmation path used by the application
    layer; neither is ever derived from parsing the file.
    """

    source = Path(str(path)).expanduser()
    if not source.is_file():
        raise ValueError("铁塔数据源必须是存在的具体文件")
    rows = _read_rows(source)
    if not rows:
        raise ValueError("铁塔数据源中没有找到可识别的数据行（需要同时具备站址编码与经纬度列）")

    items, invalid = _rows_to_items(rows, source_file=str(source))
    if invalid and not items:
        raise ValueError("铁塔数据源的每一行都缺少有效标识或坐标")

    identifiers = [item["tower_id"] for item in items]
    if len(identifiers) != len(set(identifiers)):
        seen, repeated = set(), []
        for identifier in identifiers:
            if identifier in seen and identifier not in repeated:
                repeated.append(identifier)
            seen.add(identifier)
        raise ValueError(f"铁塔站址编码重复：{', '.join(repeated[:5])}")

    record = empty_towers()
    record["status"] = "passed"
    record["source"] = {"type": "file", "path": str(source), "format": source.suffix.lower().lstrip(".")}
    record["data_source"] = str(source)
    record["count"] = len(items)
    record["items"] = [normalize_tower(item, index, source=item.get("source"))
                       for index, item in enumerate(items)]
    record["skipped"] = invalid

    representation = _representation_crs(source.suffix.lower())
    if crs is not None:
        normalized = normalize_crs_record(crs, default=empty_crs_record())
    else:
        normalized = empty_crs_record(note=(
            "源表未声明 CRS；坐标仅按源数值展示，不得自动假定 WGS84/CGCS2000。"
        ))
        if source_crs:
            normalized["source_crs"] = {
                "value": str(source_crs),
                "status": "confirmed" if crs_confirmed else "pending_confirmation",
                "axis_order": "lon_lat",
                "confirmed": bool(crs_confirmed),
                "source": deepcopy(crs_source) or {"type": "user_confirmation"},
                "evidence": deepcopy(crs_evidence or []),
            }
    normalized["representation_crs"] = representation
    record["crs"] = normalized

    duplicates = duplicate_report(record["items"])
    record["metadata"]["duplicate_coordinate_groups"] = duplicates
    record["metadata"]["duplicate_coordinate_group_count"] = len(duplicates)
    record["metadata"]["source_file"] = source.name
    record["metadata"]["source_path"] = str(source)
    record["metadata"]["invalid_row_count"] = len(invalid)
    record["metadata"]["preserved_source_fields"] = list(SOURCE_ATTRIBUTE_FIELDS)
    record["metadata"]["schema_version"] = TOWER_SCHEMA_VERSION
    record["metadata"]["collection_id"] = COLLECTION_ID
    record["warnings"] = (
        ([f"invalid_rows_skipped:{len(invalid)}"] if invalid else [])
        + ([f"near_duplicate_coordinate_groups:{len(duplicates)}"] if duplicates else [])
    )
    return record
