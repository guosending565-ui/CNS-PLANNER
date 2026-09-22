"""Application boundary for reference-only facts, previews and CRS confirmation."""

from copy import deepcopy
import csv
from hashlib import sha256
import json
from pathlib import Path

from ..algorithms.coverage.v1 import distance_m
from ..domain.source_audit import provenance_record, sha256_file, source_manifest
from ..domain.geometry_health import inspect_geojson_geometries
from ..reference_data import (
    declared_source_crs, load_equipment_reference_catalog, load_reference_landing_sites,
    load_reference_routes, load_towers,
)

#: 人工裁定的铁塔源坐标系（2026 舟山通信铁塔报送数据接入裁定）。
#: 这是**人工确认**，不是从文件内容解析出的证据；文件本身未声明任何 CRS。
TOWERS_CONFIRMED_SOURCE_CRS = "EPSG:4490"
TOWERS_CONFIRMED_SOURCE_CRS_NOTE = "人工裁定：CGCS2000 / EPSG:4490"
TOWERS_PERMISSION_EVIDENCE = (
    {"type": "usage_permission", "value": "confirmed",
     "note": "人工裁定：允许进入项目数据库/项目数据目录"},
    {"type": "publication_permission", "value": "confirmed",
     "note": "人工裁定：允许用于科研分析、论文、报告和地图成果展示"},
    {"type": "classification", "value": "non_sensitive",
     "note": "人工裁定：非涉密资料"},
)


class ReferenceDataService:
    def __init__(self, session, route_service, snapshot):
        self.session = session
        self.route_service = route_service
        self.snapshot = snapshot

    def ensure_equipment_catalog(self):
        current = self.session.state.get("equipment_reference_catalog") or {}
        if current.get("status") != "passed":
            self.session.state["equipment_reference_catalog"] = load_equipment_reference_catalog()

    def landing_sites_snapshot(self):
        return deepcopy(self.session.state.get("reference_landing_sites") or {})

    def equipment_catalog_snapshot(self):
        return deepcopy(self.session.state.get("equipment_reference_catalog") or {})

    def routes_snapshot(self):
        return deepcopy(self.session.state.get("reference_routes") or {})

    def import_landing_sites(self, path, save=True):
        previous = self.session.state.get("reference_landing_sites") or {}
        result = load_reference_landing_sites(path)
        self._migrate_landing_site_references(previous, result)
        self.session.state["reference_landing_sites"] = result
        valid = [item["coordinate"] for item in result.get("items") or [] if item.get("coordinate")]
        self._store_audit("reference_landing_sites", path, {
            "row_count": result.get("count"), "feature_count": result.get("count"),
            "extent": ([min(p[0] for p in valid), min(p[1] for p in valid),
                        max(p[0] for p in valid), max(p[1] for p in valid)] if valid else None),
            "geometry_health": {
                "status": "warning" if (result.get("metadata") or {}).get("quality_counts", {}).get("invalid") else "passed",
                "feature_count": result.get("count"), "null": 0, "empty": 0,
                "invalid": (result.get("metadata") or {}).get("quality_counts", {}).get("invalid", 0),
                "unsupported": 0,
            },
        })
        if save:
            self.session.save()
        return self.snapshot()

    # ------------------------------------------------------------------ towers
    def towers_snapshot(self):
        return deepcopy(self.session.state.get("towers") or {})

    def import_towers(self, path, save=True, *, source_crs=TOWERS_CONFIRMED_SOURCE_CRS):
        """导入真实通信铁塔**站址位置**（只读原始 Excel）。

        边界（本阶段刻意不做的事）：

        * 不写 ``grid_attributes["towers"]`` —— 因此不进入 Risk V1/V2、planning
          exposure、Theta* V2 objective 或任何失效链取值计算；
        * 不生成 coverage / 不把站址转成 C/N/S existing site / candidate site；
        * 原始文件始终以只读方式解析，绝不回写。

        ``source_crs`` 是**人工裁定**的确认值，不是解析得到的证据。
        """

        source_crs = str(source_crs or TOWERS_CONFIRMED_SOURCE_CRS)
        result = load_towers(
            path,
            source_crs=source_crs,
            crs_confirmed=True,
            crs_source={"type": "user_confirmation", "origin": "human_adjudication"},
            crs_evidence=[
                {"type": "user_supplied", "value": source_crs,
                 "note": TOWERS_CONFIRMED_SOURCE_CRS_NOTE},
                *[deepcopy(item) for item in TOWERS_PERMISSION_EVIDENCE],
            ],
        )
        self.session.state["towers"] = result
        items = result.get("items") or []
        coordinates = [(item["longitude"], item["latitude"]) for item in items]
        resolved = (result.get("crs") or {}).get("source_crs") or {}
        self._store_audit("towers", path, {
            "row_count": result.get("count"),
            "feature_count": result.get("count"),
            "extent": ([min(point[0] for point in coordinates), min(point[1] for point in coordinates),
                        max(point[0] for point in coordinates), max(point[1] for point in coordinates)]
                       if coordinates else None),
            "declared_crs": None,
            "confirmed_crs": resolved.get("value") if resolved.get("confirmed") else None,
            "geometry_health": {
                "status": "passed" if not result.get("skipped") else "warning",
                "feature_count": result.get("count"), "null": 0, "empty": 0,
                "invalid": len(result.get("skipped") or []), "unsupported": 0,
            },
            "evidence": [
                {"type": "crs_confirmation", "value": resolved.get("value"),
                 "note": TOWERS_CONFIRMED_SOURCE_CRS_NOTE},
                *[deepcopy(item) for item in TOWERS_PERMISSION_EVIDENCE],
            ],
            "provenance": provenance_record(
                source_entity="local_map_sources:towers",
                processing_activity="tower_site_reference_import",
                derived_entity="towers",
                derived_from=["local_map_sources:towers"],
                method="read_only_table_parse_with_human_confirmed_source_crs",
                note=(
                    "仅登记真实站址位置；CGCS2000/EPSG:4490 为人工裁定（源文件未声明 CRS）；"
                    "不生成 CNS 覆盖、不转 C/N/S existing site、不进入规划与风险数学。"
                ),
            ),
        }, verify_content=True)
        if save:
            self.session.save()
        return self.snapshot()

    def restore_towers_from_source(self, path, save=False):
        """Configured-source bootstrap: import once when the project has no towers yet."""

        if (self.session.state.get("towers") or {}).get("items"):
            return self.towers_snapshot()
        return self.import_towers(path, save=save)

    def _migrate_landing_site_references(self, previous, current):
        old_by_id = {
            item.get("reference_site_id"): item
            for item in previous.get("items") or [] if item.get("reference_site_id")
        }
        new_items = current.get("items") or []
        for node in self.session.state.get("nodes") or []:
            old_id = node.get("reference_site_id")
            old = old_by_id.get(old_id)
            if not old:
                continue
            candidates = [item for item in new_items if self._same_landing_site(old, item)]
            if len(candidates) != 1 or candidates[0].get("reference_site_id") == old_id:
                continue
            new_id = candidates[0]["reference_site_id"]
            node["reference_site_id"] = new_id
            provenance = node.setdefault("provenance", {})
            provenance.setdefault("legacy_reference_site_ids", []).append(old_id)
            provenance["reference_site_id"] = new_id

    @staticmethod
    def _same_landing_site(left, right):
        left_coordinate, right_coordinate = left.get("coordinate"), right.get("coordinate")
        return (
            str(left.get("name") or "").strip().casefold() == str(right.get("name") or "").strip().casefold()
            and left_coordinate == right_coordinate
        )

    def import_routes(self, path, save=True):
        result = load_reference_routes(path)
        result["data_source"] = {
            **dict(result.get("data_source") or {}),
            "imported": True, "imported_from": "direct_source_import",
        }
        self.session.state["reference_routes"] = result
        self._store_audit("reference_routes", path, {
            "row_count": result.get("point_count", 0) + int((result.get("metadata") or {}).get("invalid_row_count") or 0),
            "feature_count": result.get("count"),
            "geometry_health": {
                "status": "warning" if (result.get("metadata") or {}).get("invalid_row_count") else "passed",
                "feature_count": result.get("count"), "null": 0, "empty": 0,
                "invalid": int((result.get("metadata") or {}).get("invalid_row_count") or 0),
                "unsupported": 0,
            },
        })
        # 源文件自带 source_crs + crs_confirmed=true 时自动确认，不再要求用户重复确认。
        self._auto_confirm_declared_crs("reference_routes", save=False)
        if save:
            self.session.save()
        return self.snapshot()

    def preview_routes(self, path, conversion=None, auto_confirm=True):
        source = Path(str(path or "")).expanduser()
        if not source.is_file():
            raise ValueError("参考航线源路径必须是存在的具体文件")
        if source.suffix.lower() == ".et":
            preview = {
                "status": "requires_xlsx_or_csv_conversion",
                "reason": "requires_xlsx_or_csv_conversion", "route_count": 0,
                "point_count": 0, "warnings": ["requires_xlsx_or_csv_conversion"],
                "source_audit": source_manifest("reference_routes", source),
            }
            self.session.state["reference_route_import_preview"] = preview
            self.session.save()
            return self.snapshot()
        parsed = load_reference_routes(source)
        columns = self._source_columns(source)
        coordinates = [
            item.get("coordinate") for item in parsed.get("points") or []
            if isinstance(item.get("coordinate"), list)
        ]
        gaps = []
        conflicts = self._duplicate_route_number_conflicts(source)
        for route in parsed.get("items") or []:
            sequences = sorted({int(item["sequence"]) for item in route.get("ordered_points") or []
                                if isinstance(item.get("sequence"), (int, float))})
            missing = sorted(set(range(sequences[0], sequences[-1] + 1)) - set(sequences)) if sequences else []
            if missing:
                gaps.append({"route_number": route["route_number"], "missing_sequences": missing})
        invalid = int((parsed.get("metadata") or {}).get("invalid_row_count") or 0)
        bounds = (
            [min(p[0] for p in coordinates), min(p[1] for p in coordinates),
             max(p[0] for p in coordinates), max(p[1] for p in coordinates)]
            if coordinates else None
        )
        preview_basis = {
            "file_name": source.name, "size": source.stat().st_size,
            "mtime_ns": source.stat().st_mtime_ns, "routes": parsed.get("items"),
        }
        preview_id = "RIP-" + sha256(json.dumps(
            preview_basis, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        ).encode()).hexdigest()[:16].upper()
        geometry_health = {
            "status": "passed" if not invalid else "warning",
            "feature_count": parsed.get("route_count", parsed.get("count", 0)),
            "null": 0, "empty": 0, "invalid": invalid, "unsupported": 0,
            "extent": bounds,
        }
        if source.suffix.lower() in (".geojson", ".json"):
            document = json.loads(source.read_text(encoding="utf-8-sig"))
            geometry_health = inspect_geojson_geometries(
                document.get("features") or [], {"Point", "LineString"},
            )
        preview = {
            "status": (
                "ready_for_confirmation"
                if parsed.get("items") and geometry_health.get("status") != "blocked"
                else "blocked"
            ),
            "preview_id": preview_id,
            "file_name": source.name,
            "format": source.suffix.lower().lstrip("."),
            "route_count": parsed.get("count", 0),
            "point_count": parsed.get("point_count", 0),
            "valid_coordinate_count": len(coordinates),
            "invalid_coordinate_count": invalid + sum(
                1 for item in parsed.get("points") or [] if not item.get("coordinate")
            ),
            "duplicate_route_numbers": sorted(conflicts),
            "sequence_gaps": gaps,
            "columns": columns,
            "bounds": bounds,
            "warnings": sorted(set((parsed.get("warnings") or []) +
                                   (["sequence_gaps"] if gaps else []) +
                                   (["duplicate_route_number_conflict"] if conflicts else []) +
                                   (["invalid_geometry_fail_closed"] if geometry_health.get("status") == "blocked" else []))),
            "source_audit": source_manifest("reference_routes", source, details={
                "schema": {"columns": columns}, "row_count": len(coordinates) + invalid,
                "feature_count": parsed.get("count", 0), "extent": bounds,
                "geometry_health": geometry_health,
            }),
            "confirmation_required": True,
            "replacement_semantics": "replace_reference_routes_only_after_user_confirmation",
            "conversion_record": deepcopy(conversion) if conversion else None,
            # 源文件自身的坐标系声明：为真时无需用户再次确认。
            "declared_source_crs": deepcopy(parsed.get("declared_source_crs")),
        }
        self.session.state["reference_route_import_preview"] = preview
        self._store_audit("reference_routes", source, {
            "schema": {"columns": columns}, "row_count": len(coordinates) + invalid,
            "feature_count": parsed.get("count", 0), "extent": bounds,
            "geometry_health": geometry_health,
        })
        # 源文件已自带"坐标系已确认"声明时，预览即完成导入：正式 reference_routes
        # 对象立即生成，项目里不再只留下一个孤立的 preview。
        if auto_confirm and self._auto_import_declared_routes(source, preview):
            return self.snapshot()
        self.session.save()
        return self.snapshot()

    def confirm_routes_import(self, path, preview_id):
        preview = self.session.state.get("reference_route_import_preview") or {}
        if not preview_id or preview.get("preview_id") != preview_id:
            raise ValueError("导入预览已变化或不存在，请重新预览")
        if preview.get("status") != "ready_for_confirmation":
            raise ValueError("导入预览未通过，不能替换 reference_routes")
        source = Path(str(path or "")).expanduser()
        current = self.preview_routes(
            source, preview.get("conversion_record"), auto_confirm=False,
        )["reference_route_import_preview"]
        if current.get("preview_id") != preview_id:
            raise ValueError("源文件在预览后发生变化，请重新确认")
        self._store_confirmed_routes(source, preview_id, current)
        return self.snapshot()

    def _store_confirmed_routes(self, source, preview_id, preview):
        """把已确认的预览落成正式 ``reference_routes`` 业务航线对象。"""

        result = load_reference_routes(source)
        result["import_preview_id"] = preview_id
        result["provenance"] = provenance_record(
            source_entity=((preview.get("source_audit") or {}).get("source_id") or "SRC-REFROUTES"),
            processing_activity="confirmed_reference_route_import",
            derived_entity=result["collection_id"], method="CSV/XLSX/GeoJSON parser",
            note="用户确认预览后替换 reference_routes；未进入规划结果。",
        )
        result["provenance"]["conversion_record"] = deepcopy(preview.get("conversion_record"))
        result["data_source"] = {
            **dict(result.get("data_source") or {}),
            "imported": True,
            "imported_from": preview.get("import_origin") or "user_confirmed_preview",
        }
        self.session.state["reference_routes"] = result
        self.session.state["reference_route_import_preview"] = None
        self._auto_confirm_declared_crs("reference_routes", save=False)
        self.session.save()
        return result

    def _auto_import_declared_routes(self, source, preview):
        """源文件自带已确认坐标系时直接生成正式业务航线。

        只在本项目还没有正式 ``reference_routes`` 对象时生效；已有对象时仍然
        由用户显式确认替换，避免一次预览就覆盖既有成果。
        """

        if preview.get("status") != "ready_for_confirmation":
            return False
        declaration = preview.get("declared_source_crs") or {}
        if declaration.get("confirmed") is not True or not declaration.get("value"):
            return False
        if (self.session.state.get("reference_routes") or {}).get("items"):
            return False
        preview["import_origin"] = "source_declared_crs_auto_import"
        self._store_confirmed_routes(source, preview.get("preview_id"), preview)
        return True

    def restore_routes_from_source(self, path, save=False):
        """项目打开时按已保存的数据源恢复正式业务航线。

        前提是源文件**自己声明了已确认的坐标系**（``source_crs`` 且
        ``crs_confirmed=true``）；只做转述与校验，绝不推断。缺少该声明时保持既有
        契约：仍需用户显式预览并确认，配置一个文件本身永不替换 reference_routes。
        返回当前快照，任何失败都保持项目原有内容不变。
        """

        source = Path(str(path or "")).expanduser()
        collection = self.session.state.get("reference_routes") or {}
        if not source.is_file():
            return self.snapshot()
        declaration = collection.get("declared_source_crs")
        if not isinstance(declaration, dict):
            declaration = declared_source_crs(source) or {}
        if declaration.get("confirmed") is not True or not declaration.get("value"):
            return self.snapshot()
        if not collection.get("items"):
            parsed = load_reference_routes(source)
            if parsed.get("items"):
                parsed["declared_source_crs"] = declaration
                parsed["data_source"] = {
                    **dict(parsed.get("data_source") or {}),
                    "imported": True, "imported_from": "project_open_restore",
                }
                self.session.state["reference_routes"] = parsed
                self.session.state["reference_route_import_preview"] = None
                collection = parsed
        if collection.get("items"):
            if not isinstance(collection.get("declared_source_crs"), dict):
                collection["declared_source_crs"] = declaration
            self._auto_confirm_declared_crs("reference_routes", save=False)
        if save:
            self.session.save()
        return self.snapshot()

    def _auto_confirm_declared_crs(self, role, save=False):
        """按源文件声明自动确认坐标系，免去用户重复确认。

        仅当源文件明确写了 ``crs_confirmed=true`` 且取值能通过 pyproj 校验时
        才升级为已确认记录；任何异常都只留下 warning，不阻断导入。
        """

        collection = self.session.state.get(role) or {}
        declaration = collection.get("declared_source_crs") or {}
        value = str(declaration.get("value") or "").strip()
        if declaration.get("confirmed") is not True or not value:
            return False
        current = ((collection.get("crs") or {}).get("source_crs") or {})
        if current.get("confirmed") is True and current.get("value") == value:
            return False
        payload = {
            "value": value,
            "source": {
                "type": "source_file_column",
                "file": (collection.get("source") or {}).get("file"),
                "columns": deepcopy(declaration.get("columns") or ["source_crs", "crs_confirmed"]),
            },
            "evidence": deepcopy(declaration.get("evidence") or [{
                "type": "source_file_declaration", "value": value,
                "note": "源文件 source_crs 列声明该坐标系且 crs_confirmed 为真。",
            }]),
        }
        try:
            self.confirm_crs(role, payload, save=save)
        except ValueError as exc:
            collection.setdefault("warnings", []).append("declared_crs_auto_confirm_failed")
            collection.setdefault("metadata", {})["declared_crs_auto_confirm_error"] = str(exc)
            return False
        return True

    def confirm_crs(self, role, payload, save=True):
        if role not in ("reference_landing_sites", "reference_routes"):
            raise ValueError("只支持 reference_landing_sites/reference_routes CRS 确认")
        if not isinstance(payload, dict):
            raise ValueError("CRS 确认请求必须是对象")
        value = str(payload.get("value") or "").strip()
        source = payload.get("source")
        evidence = payload.get("evidence")
        if not value or source in (None, "") or not isinstance(evidence, list) or not evidence:
            raise ValueError("CRS 确认必须包含 value/source/evidence")
        try:
            from pyproj import CRS, Transformer
            parsed = CRS.from_user_input(value)
        except ImportError as exc:
            raise ValueError("CRS 验证需要安装 geo extra 中的 pyproj") from exc
        except Exception as exc:
            raise ValueError(f"CRS 无效：{value}") from exc
        canonical = parsed.to_authority()
        canonical_value = f"{canonical[0]}:{canonical[1]}" if canonical else parsed.to_string()
        axis_order = " / ".join(
            f"{axis.abbrev or axis.name}:{axis.direction}" for axis in parsed.axis_info
        ) or "unspecified"
        collection = self.session.state.get(role) or {}
        if not collection.get("items"):
            raise ValueError("参考数据为空，无法确认 CRS")
        previous_value = (((collection.get("crs") or {}).get("source_crs") or {}).get("value"))
        transformer = Transformer.from_crs(parsed, "OGC:CRS84", always_xy=True)

        def transformed(point):
            if not isinstance(point, list) or len(point) < 2:
                return None
            lon, lat = transformer.transform(float(point[0]), float(point[1]))
            if not (-180 <= lon <= 180 and -90 <= lat <= 90):
                raise ValueError("CRS 转换后的坐标超出经纬度范围")
            return [float(lon), float(lat)]

        crs_record = deepcopy(collection.get("crs") or {})
        crs_record["source_crs"] = {
            "value": canonical_value, "status": "confirmed", "confirmed": True,
            "axis_order": axis_order, "source": deepcopy(source),
            "evidence": deepcopy(evidence),
        }
        crs_record["representation_crs"] = {
            "value": "OGC:CRS84", "status": "declared", "axis_order": "lon_lat",
            "declared_by_format": False,
            "source": {"type": "coordinate_transformation"},
            "evidence": [{"type": "pyproj_transformer", "always_xy": True}],
        }
        collection["crs"] = crs_record
        if role == "reference_landing_sites":
            for item in collection.get("items") or []:
                source_coordinate = item.setdefault("source_coordinate", deepcopy(item.get("coordinate")))
                item["coordinate"] = transformed(source_coordinate)
                item["crs"] = deepcopy(crs_record)
                item["crs_status"] = "confirmed"
                item["metric_measurement_status"] = "enabled"
        else:
            point_by_id = {}
            for point in collection.get("points") or []:
                source_coordinate = point.setdefault("source_coordinate", deepcopy(point.get("coordinate")))
                point["coordinate"] = transformed(source_coordinate)
                point["crs"] = deepcopy(crs_record)
                point["crs_status"] = "confirmed"
                point_by_id[point.get("reference_route_point_id")] = point
            for route in collection.get("items") or []:
                original = route.setdefault("source_numeric_path", deepcopy(route.get("path") or []))
                route["path"] = [transformed(point) for point in original if point]
                route["ordered_points"] = [
                    deepcopy(point_by_id[point_id]) for point_id in route.get("ordered_point_ids") or []
                    if point_id in point_by_id
                ]
                route["length_m"] = sum(
                    distance_m(left, right) for left, right in zip(route["path"], route["path"][1:])
                )
                route["length_status"] = "passed"
                route["length_unresolved_reason"] = None
                route["metric_geometry_similarity_available"] = True
                route["crs"] = deepcopy(crs_record)
                route["crs_status"] = "confirmed"
        collection.setdefault("metadata", {})["metric_measurement_status"] = "enabled"
        collection["metadata"]["transformation"] = {
            "source_crs": canonical_value, "target_crs": "OGC:CRS84",
            "method": "pyproj.Transformer.from_crs", "always_xy": True,
        }
        # 数据源摘要随确认结果一起刷新，项目保存后即可据此判断"已导入且坐标系已确认"。
        source_state = collection.get("data_source")
        if isinstance(source_state, dict):
            source_state["crs"] = canonical_value
            source_state["crs_confirmed"] = True
            source_state["imported"] = bool(collection.get("items")) or source_state.get("imported") is True
            source_state["route_count"] = len(collection.get("items") or [])
            source_state["point_count"] = collection.get("point_count", len(collection.get("points") or []))
        audit = ((self.session.state.get("source_audits") or {}).get("items") or {}).get(role)
        collection["crs_confirmation_provenance"] = provenance_record(
            source_entity=(audit or {}).get("source_id") or f"{role}:configured_source",
            processing_activity="manual_crs_confirmation_and_representation_transform",
            derived_entity=collection.get("collection_id") or role,
            method="pyproj.Transformer.from_crs(source, OGC:CRS84, always_xy=True)",
            note="CRS 值由用户提供证据并确认；系统仅校验和转换，不推测。",
        )
        collection["metric_comparison_status"] = "ready"
        collection["warnings"] = [
            warning for warning in collection.get("warnings") or []
            if "crs_pending" not in warning and "unresolved_source_crs" not in warning
        ]
        if audit is not None:
            audit["confirmed_crs"] = deepcopy(crs_record["source_crs"])
        if previous_value and previous_value != canonical_value:
            links = self.session.state.get("reference_route_links") or {}
            if role == "reference_routes" and links.get("items"):
                links["status"] = "stale_crs_changed"
                for item in links["items"]:
                    item["current_applicability"] = "stale_crs_changed"
                    item["confirmed"] = False
            for experiment in (self.session.state.get("route_planning_experiments") or {}).get("records") or []:
                experiment["reference_comparison_status"] = "stale_crs_changed"
        if save:
            self.session.save()
        return self.snapshot()

    @staticmethod
    def _source_columns(source):
        suffix = source.suffix.lower()
        if suffix == ".csv":
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    with source.open("r", encoding=encoding, newline="") as handle:
                        return list(csv.DictReader(handle).fieldnames or [])
                except UnicodeDecodeError:
                    continue
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                columns = []
                for sheet in workbook.worksheets:
                    for values in sheet.iter_rows(values_only=True):
                        current = [str(value) for value in values if value not in (None, "")]
                        if current:
                            columns.extend(value for value in current if value not in columns)
                            break
                return columns
            finally:
                workbook.close()
        if suffix in (".geojson", ".json"):
            document = json.loads(source.read_text(encoding="utf-8-sig"))
            return sorted({
                key for feature in document.get("features") or []
                for key in ((feature or {}).get("properties") or {})
            })
        return []

    def _store_audit(self, role, path, details, *, verify_content=False):
        audits = self.session.state.setdefault("source_audits", {
            "status": "not_calculated", "schema_version": 1, "count": 0, "items": {},
        })
        previous = audits.setdefault("items", {}).get(role)
        verified_sha256 = None
        if verify_content:
            candidate = Path(str(path or "")).expanduser()
            if candidate.is_file():
                verified_sha256 = sha256_file(candidate)
        audits["items"][role] = source_manifest(
            role, path, previous=previous, details=details,
            verified_sha256=verified_sha256,
        )
        audits["count"] = len(audits["items"])
        audits["status"] = "passed"

    @staticmethod
    def _duplicate_route_number_conflicts(source):
        from ..reference_data.routes import (
            _field, _geojson_rows, _read_csv, _read_xlsx,
        )

        suffix = source.suffix.lower()
        if suffix == ".csv":
            rows = [(source.stem, row, values) for row, values in _read_csv(source)]
        elif suffix == ".xlsx":
            rows = _read_xlsx(source)
        elif suffix in (".geojson", ".json"):
            rows = [(source.stem, row, values) for row, values in _geojson_rows(source)]
        else:
            return []
        identities = {}
        last_number = None
        for _sheet, _row, values in rows:
            number = _field(values, "route_number")
            if number not in (None, ""):
                last_number = str(number).strip()
            elif last_number:
                number = last_number
            if number in (None, ""):
                continue
            identity = (
                str(_field(values, "route_name") or "").strip().casefold(),
                str(_field(values, "category") or "").strip().casefold(),
            )
            if any(identity):
                identities.setdefault(str(number).strip(), set()).add(identity)
        return sorted(number for number, values in identities.items() if len(values) > 1)

    def add_landing_site_to_project(self, reference_site_id):
        collection = self.session.state.get("reference_landing_sites") or {}
        site = next(
            (item for item in collection.get("items") or [] if item.get("reference_site_id") == reference_site_id),
            None,
        )
        if not site:
            raise ValueError("参考起降点不存在")
        if site.get("quality") == "invalid" or not site.get("coordinate"):
            raise ValueError("参考起降点坐标无效，不能加入项目")
        return self.route_service.add_reference_site(site)
